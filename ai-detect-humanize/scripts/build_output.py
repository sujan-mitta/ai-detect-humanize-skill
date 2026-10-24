#!/usr/bin/env python3
"""Join humanized chunks back into one document.

Usage: python build_output.py CHUNK_DIR --format docx|md|txt -o OUTPUT

Uses chunk_NNN.human.txt for each chunk in manifest order. If a humanized
chunk is missing, the original chunk is used and a warning is printed, so
nothing is silently dropped.
"""
import argparse
import json
import re
import sys
from pathlib import Path


def collect(chunk_dir):
    chunk_dir = Path(chunk_dir)
    manifest = json.loads((chunk_dir / "manifest.json").read_text(encoding="utf-8"))
    parts, missing = [], []
    for c in manifest["chunks"]:
        human = chunk_dir / c["humanized_file"]
        if human.exists():
            parts.append(human.read_text(encoding="utf-8").strip())
        else:
            missing.append(c["file"])
            parts.append((chunk_dir / c["file"]).read_text(encoding="utf-8").strip())
    if missing:
        print(f"WARNING: {len(missing)} chunk(s) not humanized, originals used: "
              + ", ".join(missing), file=sys.stderr)
    return "\n\n".join(parts)


def add_runs(paragraph, text):
    """Render **bold** and *italic* markdown inline."""
    for piece in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            paragraph.add_run(piece[1:-1]).italic = True
        elif piece:
            paragraph.add_run(piece)


def write_docx(text, output):
    try:
        import docx
    except ImportError:
        sys.exit("python-docx is not installed: pip install python-docx")
    document = docx.Document()
    for block in re.split(r"\n\s*\n", text):
        lines = block.splitlines()
        if not lines:
            continue
        if all(l.strip().startswith("|") for l in lines):
            rows = [[c.strip() for c in l.strip().strip("|").split("|")] for l in lines
                    if not re.match(r"^\|?\s*:?-+", l.strip().strip("|"))]
            table = document.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    table.cell(r, c).text = cell
            continue
        for line in lines if any(re.match(r"^(#{1,6}\s|[-*]\s|\d+[.)]\s)", l) for l in lines) else [" ".join(lines)]:
            m = re.match(r"^(#{1,6})\s+(.*)", line)
            if m:
                document.add_heading(m.group(2).strip(), level=min(len(m.group(1)), 4))
            elif re.match(r"^[-*]\s", line):
                add_runs(document.add_paragraph(style="List Bullet"), line[2:].strip())
            elif re.match(r"^\d+[.)]\s", line):
                add_runs(document.add_paragraph(style="List Number"), re.sub(r"^\d+[.)]\s", "", line).strip())
            else:
                add_runs(document.add_paragraph(), line.strip())
    document.save(output)


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("chunk_dir")
    ap.add_argument("--format", choices=["docx", "md", "txt"], default="md")
    ap.add_argument("-o", "--output", required=True)
    args = ap.parse_args()

    text = collect(args.chunk_dir)
    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
    if args.format == "docx":
        write_docx(text, args.output)
    else:
        if args.format == "txt":
            text = re.sub(r"^#{1,6}\s+", "", text, flags=re.M)
        Path(args.output).write_text(text + "\n", encoding="utf-8")
    print(f"Wrote {len(text.split())} words to {args.output}")


if __name__ == "__main__":
    main()
