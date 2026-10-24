#!/usr/bin/env python3
"""Extract plain text from PDF, DOCX, TXT or MD files.

Usage: python extract_text.py INPUT [-o OUTPUT]

DOCX headings become markdown '#' lines so structure survives chunking and
rebuilding. PDF pages are joined with blank lines; hyphenated line breaks
are repaired and hard-wrapped lines are re-flowed into paragraphs.
"""
import argparse
import re
import sys
from pathlib import Path


def extract_pdf(path):
    try:
        from pypdf import PdfReader
    except ImportError:
        sys.exit("pypdf is not installed: pip install pypdf")
    reader = PdfReader(str(path))
    pages = [(page.extract_text() or "") for page in reader.pages]
    text = "\n\n".join(pages)
    if len(text.strip()) < 50 * max(1, len(pages)) * 0.2:
        print("WARNING: very little text extracted - this PDF may be scanned "
              "and need OCR.", file=sys.stderr)
    return reflow(text)


def reflow(text):
    # Join words hyphenated across line breaks: "inter-\nnational" -> "international"
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    paragraphs = re.split(r"\n\s*\n", text)
    out = []
    for para in paragraphs:
        lines = [l.strip() for l in para.splitlines() if l.strip()]
        if not lines:
            continue
        # Keep short lines (likely headings / list items) on their own
        merged, buf = [], []
        for line in lines:
            if re.match(r"^([-*•]|\d+[.)])\s", line) or (len(line) < 60 and not line.endswith((".", ",", ";"))):
                if buf:
                    merged.append(" ".join(buf))
                    buf = []
                merged.append(line)
            else:
                buf.append(line)
        if buf:
            merged.append(" ".join(buf))
        out.append("\n".join(merged))
    return "\n\n".join(out)


def extract_docx(path):
    try:
        import docx
    except ImportError:
        sys.exit("python-docx is not installed: pip install python-docx")
    document = docx.Document(str(path))
    parts = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        m = re.match(r"heading (\d)", style)
        if m:
            parts.append("#" * int(m.group(1)) + " " + text)
        elif style == "title":
            parts.append("# " + text)
        elif "list" in style:
            parts.append("- " + text)
        else:
            parts.append(text)
    for table in document.tables:
        rows = ["| " + " | ".join(c.text.strip() for c in row.cells) + " |" for row in table.rows]
        if rows:
            parts.append("\n".join(rows))
    return "\n\n".join(parts)


def extract(path):
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix == ".docx":
        return extract_docx(path)
    return path.read_text(encoding="utf-8", errors="replace")


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("input")
    ap.add_argument("-o", "--output")
    args = ap.parse_args()
    text = extract(args.input)
    if args.output:
        Path(args.output).parent.mkdir(parents=True, exist_ok=True)
        Path(args.output).write_text(text, encoding="utf-8")
        print(f"Wrote {len(text.split())} words to {args.output}")
    else:
        sys.stdout.reconfigure(encoding="utf-8")
        print(text)


if __name__ == "__main__":
    main()
